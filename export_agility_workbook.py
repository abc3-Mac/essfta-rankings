"""Agility results workbook for a side-by-side check with the club's statistician (Linda): the
rankings, plus every run each ranking counted, with the speed written as Excel formulas
(yards / (time - table)) so the math can be checked cell by cell.

    python3 export_agility_workbook.py 2026 Q3      -> outgoing/ESSFTA_Agility_2026_JanSep_Results.xlsx

Run build_quarter.py for the quarter first. Recomputes the legs with build_quarter.py's own
functions and asserts every average equals the built quarter (quarters/<tag>/report_data_<tag>.json)
before writing anything. First used for Linda, 21 Oct 2026.
"""
import collections, datetime, json, os, sys

REPO = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, REPO)
os.chdir(REPO)
import build_quarter as bq  # noqa: E402
from check_title_eligibility import titles  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

YEAR, Q = sys.argv[1], sys.argv[2].upper()
WINDOW = {"Q1": ("JanMar", "March 31"), "Q2": ("JanJun", "June 30"), "Q3": ("JanSep", "September 30"),
          "Q4": ("JanDec", "December 31")}[Q]
QTRS = " + ".join(f"{n}Q" for n in range(1, int(Q[1]) + 1))
OUT = os.path.join(REPO, "outgoing", f"ESSFTA_Agility_{YEAR}_{WINDOW[0]}_Results.xlsx")

runs = [r for r in bq.read_agility(bq.raw_files("Agility", YEAR, Q)) if r["place"] in bq.QUAL]
by_dog = collections.defaultdict(list)
for r in runs:
    by_dog[r["reg"]].append(r)
latest = {reg: max(rs, key=lambda r: (r["date"], r["order"])) for reg, rs in by_dog.items()}
pub = json.load(open(f"quarters/{YEAR}{Q}/report_data_{YEAR}{Q}.json"))["sports"]
pub = {(r["category"], bq.key_name(r["dog"])): r for r in pub if r["sport"] == "Agility"}
tmap, newest = titles()

# ---- the runs each ranking counted ------------------------------------------------------------
qq_rows = []      # (category, reg, day, std_run, jww_run)
leg_rows = []     # (category, reg, run, title, title_date)
for cat, std, jww in (("Excellent B", "AGEXCB", "JWWEXCB"), ("Preferred Excellent B", "AGEXCBP", "JWWEXCBP")):
    for reg, rs in by_dog.items():
        days = collections.defaultdict(lambda: {"s": [], "j": []})
        for r in rs:
            if r["cls"] == std:
                days[r["date"]]["s"].append(r)
            elif r["cls"] == jww:
                days[r["date"]]["j"].append(r)
        pairs = []
        for day, d in days.items():
            for s, j in zip(sorted(d["s"], key=bq._yps, reverse=True), sorted(d["j"], key=bq._yps, reverse=True)):
                pairs.append((day, s, j, (bq._yps(s) + bq._yps(j)) / 2))
        if len(pairs) >= 10:
            pairs.sort(key=lambda p: -p[3])
            for day, s, j, _ in pairs[:10]:
                qq_rows.append((cat, reg, day, s, j, len(pairs)))

for cat in ("Novice", "Open", "Excellent A", "Preferred Novice", "Preferred Open", "Preferred Excellent A"):
    (st, sc), (jt, jc) = bq.AG_TITLES[cat]
    for reg, rs in by_dog.items():
        rs = sorted(rs, key=lambda r: (r["date"], r["order"]))
        a = bq._title_legs(rs, sc, tmap.get((reg, st), []), int(YEAR), newest)
        b = bq._title_legs(rs, jc, tmap.get((reg, jt), []), int(YEAR), newest)
        if "UNDER 3" in (a[1], b[1]) or a[1] in bq.AG_EXCLUDE or b[1] in bq.AG_EXCLUDE:
            continue
        for legs, title, tdate in ((a[0], st, a[2]), (b[0], jt, b[2])):
            for r in legs:
                leg_rows.append((cat, reg, r, title, tdate))

# ---- check against the published quarter ------------------------------------------------------
ORDER = ["Excellent B", "Preferred Excellent B", "Novice", "Open", "Excellent A",
         "Preferred Novice", "Preferred Open", "Preferred Excellent A"]
ranking = []
for cat in ORDER:
    regs = list(dict.fromkeys([x[1] for x in qq_rows if x[0] == cat] + [x[1] for x in leg_rows if x[0] == cat]))
    for reg in regs:
        p = pub[(cat, bq.key_name(latest[reg]["dog"]))]
        if cat in ("Excellent B", "Preferred Excellent B"):
            qs = [x for x in qq_rows if x[0] == cat and x[1] == reg]
            yps = sum((bq._yps(s) + bq._yps(j)) / 2 for _, _, _, s, j, _ in qs) / 10
            assert abs(yps - float(p["score2"])) < 0.0006, (cat, p["dog"], yps, p["score2"])
            ranking.append(dict(cat=cat, reg=reg, rank=p["rank"], tie=p.get("tie") or "", dog=p["dog"],
                                owner=p["owner"], counted=qs[0][5], score=None, yps=yps))
        else:
            ls = [x[2] for x in leg_rows if x[0] == cat and x[1] == reg]
            sc = sum(float(r["score"]) for r in ls) / 6
            yps = sum(bq._yps(r) for r in ls) / 6
            assert abs(sc - float(p["score"])) < 0.006 and abs(yps - float(p["score2"])) < 0.0006, (cat, p["dog"])
            ranking.append(dict(cat=cat, reg=reg, rank=p["rank"], tie=p.get("tie") or "", dog=p["dog"],
                                owner=p["owner"], counted=6, score=sc, yps=yps))
assert len(ranking) == len(pub), (len(ranking), len(pub))
ranking.sort(key=lambda e: (ORDER.index(e["cat"]), int(e["rank"])))

# ---- workbook ---------------------------------------------------------------------------------
F = "Arial"
HDR = PatternFill("solid", fgColor="6B1E14")
HFONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
NOTE = Font(name=F, size=9, italic=True, color="595959")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
BAND = PatternFill("solid", fgColor="F7F1E6")


def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = HFONT, HDR
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def d(s):
    return datetime.date.fromisoformat(s)


wb = Workbook()

# Runs counted (Novice / Open / Excellent A and Preferred): one row per title leg
wl = wb.active
wl.title = "Title Legs"
wl["A1"] = ("Novice, Open and Excellent A (and Preferred): the 3 Standard + 3 Jumpers legs that earned this "
            "year's two titles. YPS = Course Distance / (Dog Time - Table); Table is 5 on Novice and Open Standard.")
wl["A1"].font = NOTE
lcols = ["Category", "Reg Number", "Dog Name", "Event Name", "Event Date", "Primary Class", "Score",
         "Course Distance", "Dog Time", "Table (sec)", "YPS", "Title Earned", "Date Titled"]
header(wl, 3, lcols, [20, 13, 44, 36, 11, 13, 8, 10, 9, 8, 9, 30, 11])
leg_ref = {}   # (cat, reg) -> list of rows
r0 = 4
for cat in ORDER[2:]:
    for e in [x for x in ranking if x["cat"] == cat]:
        for (c, reg, run, title, tdate) in [x for x in leg_rows if x[0] == cat and x[1] == e["reg"]]:
            vals = [cat, reg, run["dog"], run["event"], d(run["date"]), run["cls"], float(run["score"]),
                    float(run["dist"]), float(run["time"]), 5 if run["cls"] in bq.TABLE_CLASSES else 0,
                    f"=H{r0}/(I{r0}-J{r0})", title, tdate]
            for i, v in enumerate(vals, 1):
                cell = wl.cell(row=r0, column=i, value=v)
                cell.font, cell.border = BODY, THIN
            wl.cell(row=r0, column=5).number_format = "mm/dd/yyyy"
            wl.cell(row=r0, column=13).number_format = "mm/dd/yyyy"
            wl.cell(row=r0, column=11).number_format = "0.000"
            wl.cell(row=r0, column=10).font = Font(name=F, size=10, color="0000FF")
            leg_ref.setdefault((cat, e["reg"]), []).append(r0)
            r0 += 1
wl.cell(row=4, column=10).comment = Comment("Table count taken off the dog's time on Novice and Open Standard "
                                            "(AGNOVA, AGNOVB, AGNOVP, AGOPEN, AGOPENP); 0 elsewhere. Per Linda, 21 Oct 2026.", "ESSFTA")
wl.auto_filter.ref = f"A3:M{r0 - 1}"

# Excellent B QQs: the 10 fastest QQ days per dog
wq = wb.create_sheet("ExB 10 Fastest QQs")
wq["A1"] = ("Excellent B and Preferred Excellent B: each dog's 10 fastest QQ days (Standard and Jumpers Q the "
            "same day). QQ YPS = average of the two runs' YPS; the ranking is the average of the 10.")
wq["A1"].font = NOTE
qcols = ["Category", "Reg Number", "Dog Name", "Event Date", "STD Event", "STD Distance", "STD Dog Time",
         "JWW Event", "JWW Distance", "JWW Dog Time", "STD YPS", "JWW YPS", "QQ YPS"]
header(wq, 3, qcols, [20, 13, 44, 11, 30, 9, 9, 30, 9, 9, 9, 9, 9])
qq_ref = {}
r0 = 4
for cat in ORDER[:2]:
    for e in [x for x in ranking if x["cat"] == cat]:
        for (c, reg, day, s, j, n) in [x for x in qq_rows if x[0] == cat and x[1] == e["reg"]]:
            vals = [cat, reg, s["dog"], d(day), s["event"], float(s["dist"]), float(s["time"]), j["event"],
                    float(j["dist"]), float(j["time"]), f"=F{r0}/G{r0}", f"=I{r0}/J{r0}", f"=AVERAGE(K{r0}:L{r0})"]
            for i, v in enumerate(vals, 1):
                cell = wq.cell(row=r0, column=i, value=v)
                cell.font, cell.border = BODY, THIN
            wq.cell(row=r0, column=4).number_format = "mm/dd/yyyy"
            for col in (11, 12, 13):
                wq.cell(row=r0, column=col).number_format = "0.000"
            qq_ref.setdefault((cat, e["reg"]), []).append(r0)
            r0 += 1
wq.auto_filter.ref = f"A3:M{r0 - 1}"

# Rankings: averages are formulas over the two run sheets
ws = wb.create_sheet("Rankings", 0)
ws["A1"] = f"ESSFTA Agility Rankings — January 1 through {WINDOW[1]}, {YEAR}"
ws["A1"].font = Font(name=F, size=13, bold=True, color="6B1E14")
ws["A2"] = (f"Computed from the AKC results exports ({QTRS} {YEAR}). Averages are live formulas over the "
            "'Title Legs' and 'ExB 10 Fastest QQs' sheets. Ranks use the unrounded averages.")
ws["A2"].font = NOTE
rcols = ["Category", "Rank", "Dog Name", "Owner(s)", "QQs / Legs", "Avg Score", "Avg YPS"]
header(ws, 4, rcols, [22, 7, 52, 46, 10, 10, 10])
r0 = 5
prev = None
for e in ranking:
    if prev and e["cat"] != prev:
        r0 += 1
    prev = e["cat"]
    ws.cell(row=r0, column=1, value=e["cat"])
    ws.cell(row=r0, column=2, value=f"{e['rank']}{e['tie']}")
    ws.cell(row=r0, column=3, value=e["dog"])
    ws.cell(row=r0, column=4, value=e["owner"].replace(",", ", "))
    ws.cell(row=r0, column=5, value=e["counted"])
    if e["score"] is None:
        rows = qq_ref[(e["cat"], e["reg"])]
        ws.cell(row=r0, column=7, value=f"=AVERAGE('ExB 10 Fastest QQs'!M{rows[0]}:M{rows[-1]})")
        ws.cell(row=r0, column=7).font = BOLD
    else:
        rows = leg_ref[(e["cat"], e["reg"])]
        ws.cell(row=r0, column=6, value=f"=AVERAGE('Title Legs'!G{rows[0]}:G{rows[-1]})")
        ws.cell(row=r0, column=7, value=f"=AVERAGE('Title Legs'!K{rows[0]}:K{rows[-1]})")
        ws.cell(row=r0, column=6).font = BOLD
    for i in range(1, 8):
        c = ws.cell(row=r0, column=i)
        if c.font != BOLD:
            c.font = BODY
        c.border = THIN
        if e["cat"] in ORDER[1::2]:
            c.fill = BAND
    ws.cell(row=r0, column=6).number_format = "0.00"
    ws.cell(row=r0, column=7).number_format = "0.000"
    ws.cell(row=r0, column=2).alignment = Alignment(horizontal="center")
    r0 += 1
ws.cell(row=5, column=5).comment = Comment("Excellent B: QQs on record this year (10 fastest averaged). "
                                           "Other classes: the 6 title legs.", "ESSFTA")

# How it was computed
wn = wb.create_sheet("How Computed")
notes = [
    (f"ESSFTA agility rankings, {YEAR} January through {WINDOW[1]}: method", True),
    ("", False),
    ("Speed (YPS) = Course Distance / Dog Time. On Novice and Open Standard courses (AGNOVA, AGNOVB, AGNOVP, "
     "AGOPEN, AGOPENP) the 5-second table count is taken off the time first: Course Distance / (Dog Time - 5). "
     "Excellent Standard has no table; Jumpers never has.", False),
    ("Excellent B and Preferred Excellent B: at least 10 QQs (Standard and Jumpers Q on the same day). Ranked "
     "by the average QQ speed of the 10 fastest QQ days; a QQ day's speed is the average of its two runs.", False),
    ("Novice, Open, Excellent A and their Preferred classes: the average score of the 3 Standard and 3 Jumpers "
     f"legs that earned the dog's two titles. A dog is listed only when both titles are dated {YEAR} and every leg "
     f"was run in {YEAR} (checked against the AKC New Titles {YEAR} file). Equal scores are ordered by the average "
     "speed of the same six legs.", False),
    ("Ranks use the unrounded values; every speed is shown to 3 decimals, as in the PDF. Click a cell to see the full value.", False),
    (f"Source: AKC Companion Events results exports for English Springer Spaniels, {QTRS} {YEAR}; "
     f"AKC New Titles {YEAR} (titles through " + newest.strftime("%B %-d, %Y") + ").", False),
    ("Prepared by the ESSFTA rankings tools, " + datetime.date.today().strftime("%B %-d, %Y") + ".", False),
]
for i, (t, b) in enumerate(notes, 1):
    c = wn.cell(row=i, column=1, value=t)
    c.font = Font(name=F, size=12 if b else 10, bold=b, color="6B1E14" if b else "000000")
    c.alignment = Alignment(wrap_text=True, vertical="top")
wn.column_dimensions["A"].width = 120

wb.save(OUT)
print(f"wrote {OUT}: {len(ranking)} ranked dogs, {sum(len(v) for v in leg_ref.values())} title legs, "
      f"{sum(len(v) for v in qq_ref.values())} QQ days; every average matched the published quarter")
